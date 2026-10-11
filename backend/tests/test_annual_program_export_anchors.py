"""D15 (2026-10-11): the Annual Program export 500s once a year has an anchor event.

GET /api/planning/years/{id}/export.xlsx (the Main TMS "Export annual
program" button) wrote its Anchor Events sheet from a.name, a.location and
a.description; AnchorEvent has event_name and notes and no location column,
so any year with an anchor raised AttributeError -> 500. The existing export
test used a year with no anchors, so the loop never ran (found by the C5
extraction parity snapshot, which records the 500 on the seeded data).
"""
import io

from openpyxl import load_workbook

from tests.conftest import login
from tests.test_planning import _make_year


def test_export_includes_anchor_events(client):
    hdr = login(client, "ADMIN703")
    yr_id = _make_year(client, hdr)["planning_year_id"]
    r = client.post(f"/api/planning/years/{yr_id}/anchors", headers=hdr, json={
        "event_name": "=HYPERLINK(\"x\") Wing camp", "start_date": "2031-04-01",
        "end_date": "2031-04-05", "event_type": "camp", "notes": "Bring sleeping bag"})
    assert r.status_code in (200, 201), r.text
    r = client.get(f"/api/planning/years/{yr_id}/export.xlsx", headers=hdr)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Anchor Events"]
    assert [c.value for c in ws[1]] == ["Name", "Start Date", "End Date", "Type",
                                        "Importance", "Location", "Notes"]
    row = [c.value for c in ws[2]]
    assert row[0].endswith("Wing camp") and not row[0].startswith("=")   # still neutralised
    assert row[1:4] == ["2031-04-01", "2031-04-05", "camp"]
    assert row[6] == "Bring sleeping bag"
